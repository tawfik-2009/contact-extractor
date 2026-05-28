from flask import Flask, jsonify, request, send_from_directory, Response, session, redirect, url_for
from flask_session import Session
from flask_cors import CORS
import threading
import uuid
import csv
import io
import sqlite3
import json
from datetime import datetime
from scraper import ContactExtractor

app = Flask(__name__, static_folder='static')
app.config['SECRET_KEY'] = 'super-secret-key-change-me'
app.config['SESSION_TYPE'] = 'filesystem'
Session(app)
CORS(app)

DB_PATH = 'history.db'

def init_db():
    with sqlite3.connect(DB_PATH) as conn:
        c = conn.cursor()
        c.execute('''
            CREATE TABLE IF NOT EXISTS searches (
                job_id TEXT PRIMARY KEY,
                query TEXT,
                location TEXT,
                source TEXT,
                date TEXT,
                total_results INTEGER
            )
        ''')
        c.execute('''
            CREATE TABLE IF NOT EXISTS results (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                job_id TEXT,
                data TEXT,
                FOREIGN KEY (job_id) REFERENCES searches (job_id)
            )
        ''')
        conn.commit()

init_db()

# In-memory job store  { job_id: { status, progress, results, message, ... } }
jobs: dict = {}

@app.before_request
def require_login():
    # Allow static files that are needed for login (e.g. style.css)
    if request.path.startswith('/static/') and not request.path.endswith('.html'):
        return None
    
    allowed_routes = ['/login', '/api/login', '/static/login.html']
    if request.path not in allowed_routes and not session.get('logged_in'):
        if request.path.startswith('/api/'):
            return jsonify({'error': 'Unauthorized'}), 401
        return redirect('/login')

@app.route('/login')
def login_page():
    if session.get('logged_in'):
        return redirect('/')
    return send_from_directory('static', 'login.html')

@app.route('/logout')
def logout():
    session.pop('logged_in', None)
    return redirect('/login')

@app.route('/api/login', methods=['POST'])
def api_login():
    data = request.get_json(silent=True) or {}
    username = data.get('username')
    password = data.get('password')
    # Default credentials
    if username == 'admin' and password == 'admin123':
        session['logged_in'] = True
        return jsonify({'success': True})
    return jsonify({'success': False, 'message': 'Invalid credentials'}), 401

@app.route('/')
def index():
    return send_from_directory('static', 'index.html')


@app.route('/api/search', methods=['POST'])
def start_search():
    data = request.get_json(silent=True) or {}
    query = (data.get('query') or '').strip()
    location = (data.get('location') or 'السعودية').strip()
    source = data.get('source', 'google')
    max_results = min(int(data.get('max_results', 20)), 500)

    if not query and source == 'google':
        return jsonify({'error': 'يرجى إدخال نوع النشاط التجاري'}), 400

    job_id = str(uuid.uuid4())
    jobs[job_id] = {
        'status': 'running',
        'progress': 0,
        'results': [],
        'message': 'جاري التهيئة...',
        'query': query,
        'location': location,
        'source': source,
    }

    extractor = ContactExtractor(jobs[job_id], max_results=max_results)
    thread = threading.Thread(
        target=extractor.search_router,
        args=(query, location, source, job_id, DB_PATH),
        daemon=True,
    )
    thread.start()
    return jsonify({'job_id': job_id})


@app.route('/api/status/<job_id>')
def get_status(job_id):
    if job_id not in jobs:
        return jsonify({'error': 'Job not found'}), 404
    return jsonify(jobs[job_id])

@app.route('/api/pause/<job_id>', methods=['POST'])
def pause_job(job_id):
    if job_id in jobs and jobs[job_id]['status'] == 'running':
        jobs[job_id]['status'] = 'paused'
        jobs[job_id]['message'] = 'تم إيقاف البحث مؤقتاً...'
        return jsonify({'success': True})
    return jsonify({'success': False}), 400

@app.route('/api/resume/<job_id>', methods=['POST'])
def resume_job(job_id):
    if job_id in jobs and jobs[job_id]['status'] == 'paused':
        jobs[job_id]['status'] = 'running'
        jobs[job_id]['message'] = 'جاري استكمال البحث...'
        return jsonify({'success': True})
    return jsonify({'success': False}), 400

@app.route('/api/stop/<job_id>', methods=['POST'])
def stop_job(job_id):
    if job_id in jobs and jobs[job_id]['status'] in ['running', 'paused']:
        jobs[job_id]['status'] = 'stopped_early'
        jobs[job_id]['message'] = 'جاري إنهاء البحث وحفظ النتائج...'
        return jsonify({'success': True})
    return jsonify({'success': False}), 400


@app.route('/api/history')
def get_history():
    with sqlite3.connect(DB_PATH) as conn:
        conn.row_factory = sqlite3.Row
        c = conn.cursor()
        c.execute('SELECT * FROM searches ORDER BY date DESC')
        rows = c.fetchall()
        return jsonify([dict(row) for row in rows])

@app.route('/api/history/<job_id>/results')
def get_history_results(job_id):
    with sqlite3.connect(DB_PATH) as conn:
        c = conn.cursor()
        c.execute('SELECT data FROM results WHERE job_id = ?', (job_id,))
        rows = c.fetchall()
        results = [json.loads(row[0]) for row in rows]
        return jsonify(results)


# ─── Export ───────────────────────────────────────────────────────────────────

FIELD_LABELS = {
    'name':      'اسم النشاط',
    'category':  'النوع',
    'phone':     'الهاتف',
    'email':     'البريد الإلكتروني',
    'address':   'العنوان',
    'website':   'الموقع الإلكتروني',
    'rating':    'التقييم',
    'whatsapp':  'واتساب',
    'facebook':  'فيسبوك',
    'instagram': 'إنستغرام',
    'twitter':   'تويتر / X',
    'linkedin':  'لينكد إن',
    'snapchat':  'سناب شات',
    'tiktok':    'تيك توك',
    'youtube':   'يوتيوب',
    'maps_url':  'رابط الخريطة',
}
FIELDS = list(FIELD_LABELS.keys())


def get_job_results(job_id):
    if job_id in jobs:
        return jobs[job_id].get('results', [])
    with sqlite3.connect(DB_PATH) as conn:
        c = conn.cursor()
        c.execute('SELECT data FROM results WHERE job_id = ?', (job_id,))
        rows = c.fetchall()
        return [json.loads(row[0]) for row in rows]

@app.route('/api/export/<job_id>/csv')
def export_csv(job_id):
    results = get_job_results(job_id)
    if not results:
        return jsonify({'error': 'لا توجد نتائج للتصدير'}), 400

    output = io.StringIO()
    output.write('\ufeff')            # BOM → Excel reads Arabic correctly
    writer = csv.writer(output)
    writer.writerow([FIELD_LABELS[f] for f in FIELDS])
    for row in results:
        writer.writerow([row.get(f, '') for f in FIELDS])

    return Response(
        output.getvalue(),
        mimetype='text/csv; charset=utf-8-sig',
        headers={
            'Content-Disposition': f'attachment; filename=contacts_{job_id[:8]}.csv',
            'Content-Type': 'text/csv; charset=utf-8-sig',
        },
    )


@app.route('/api/export/<job_id>/excel')
def export_excel(job_id):
    results = get_job_results(job_id)
    if not results:
        return jsonify({'error': 'لا توجد نتائج للتصدير'}), 400

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'بيانات العملاء'
        ws.sheet_view.rightToLeft = True

        # Header style
        hdr_fill = PatternFill(start_color='4C1D95', end_color='4C1D95', fill_type='solid')
        hdr_font = Font(bold=True, color='FFFFFF', size=11, name='Cairo')
        hdr_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

        col_widths = {
            'name': 30, 'category': 18, 'phone': 18, 'email': 30,
            'address': 35, 'website': 28, 'rating': 10, 'whatsapp': 20,
            'facebook': 30, 'instagram': 25, 'twitter': 25, 'linkedin': 28,
            'snapchat': 22, 'tiktok': 22, 'youtube': 28, 'maps_url': 35,
        }

        for col_idx, field in enumerate(FIELDS, 1):
            cell = ws.cell(row=1, column=col_idx, value=FIELD_LABELS[field])
            cell.fill = hdr_fill
            cell.font = hdr_font
            cell.alignment = hdr_align
            ws.column_dimensions[get_column_letter(col_idx)].width = col_widths.get(field, 20)

        ws.row_dimensions[1].height = 30

        # Data rows
        alt_fill = PatternFill(start_color='F5F3FF', end_color='F5F3FF', fill_type='solid')
        data_font = Font(name='Cairo', size=10)
        data_align = Alignment(horizontal='right', vertical='center', wrap_text=False)

        for row_idx, result in enumerate(results, 2):
            fill = alt_fill if row_idx % 2 == 0 else None
            for col_idx, field in enumerate(FIELDS, 1):
                val = result.get(field, '')
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = data_font
                cell.alignment = data_align
                if fill:
                    cell.fill = fill

        ws.freeze_panes = 'A2'

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        return Response(
            buf.getvalue(),
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={
                'Content-Disposition': f'attachment; filename=contacts_{job_id[:8]}.xlsx',
            },
        )
    except ImportError:
        return jsonify({'error': 'openpyxl غير مثبت'}), 500


# ─── Entry Point ──────────────────────────────────────────────────────────────

if __name__ == '__main__':
    import sys, io
    # Ensure UTF-8 output on Windows console
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8', errors='replace')
    print('\n' + '=' * 50)
    print('  Contact-AI - استخراج بيانات العملاء')
    print('  URL: http://localhost:5000')
    print('=' * 50 + '\n')
    app.run(debug=False, port=5000, threaded=True)
